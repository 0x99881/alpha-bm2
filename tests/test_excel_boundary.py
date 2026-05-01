from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from bm2.constants import NAME_HEADER, SCORE_SHEET, TOTAL_HEADER
from .conftest import add_members, score_form


def _score_rows_for_names(store, names: set[str]) -> list[dict[str, object]]:
    return [
        row
        for row in store._context.local_db.get_score_rows(include_deleted=True)
        if row["member_name"] in names
    ]


def test_empty_database_export_is_noop(store):
    empty_store_rows = store._context.local_db.get_score_rows()

    exported = store._context.export_service.export_to_excel()

    assert exported == 0
    assert store._context.local_db.get_score_rows() == empty_store_rows


def test_export_writes_sqlite_scores_to_excel(store):
    members = add_members(store, ["Excel_A", "Excel_B"])
    store.daily_entry_service.process_submission(members, score_form(["Excel_A", "Excel_B"], 5), "2026-07-10")
    store.daily_entry_service.process_submission(members, score_form(["Excel_A", "Excel_B"], -1), "2026-07-11")

    exported = store._context.export_service.export_to_excel()

    assert exported >= 2
    workbook = load_workbook(store.workbook_path)
    try:
        sheet = workbook[SCORE_SHEET]
        values = [cell.value for row in sheet.iter_rows(values_only=False) for cell in row]
        assert "Excel_A" in values
        assert "Excel_B" in values
        assert 5 in values
        assert -1 in values
    finally:
        workbook.close()


def test_special_character_member_export(store):
    name = "Excel_特殊_#1"
    members = add_members(store, [name])
    store.daily_entry_service.process_submission(members, score_form([name], 8), "2026-07-12")

    store._context.export_service.export_to_excel()

    workbook = load_workbook(store.workbook_path)
    try:
        values = [cell.value for row in workbook[SCORE_SHEET].iter_rows(values_only=False) for cell in row]
        assert name in values
    finally:
        workbook.close()


def test_export_failure_does_not_affect_sqlite_save(store, monkeypatch):
    members = add_members(store, ["Excel_C"])

    def fail_export(_context):
        raise OSError("export failed")

    monkeypatch.setattr(store._context.excel_exporter, "export_missing_dates", fail_export)
    with pytest.raises(OSError):
        store._context.export_service.export_to_excel()

    result = store.daily_entry_service.process_submission(
        members,
        score_form(["Excel_C"], 3),
        "2026-07-13",
    )
    assert result["ok"] is True
    rows = _score_rows_for_names(store, {"Excel_C"})
    assert [(row["score_date"], row["score"]) for row in rows] == [("2026-07-13", 3)]


def test_import_history_excel_into_fresh_sqlite(store, test_workspace):
    source_members = add_members(store, ["Import_A", "Import_B"])
    store.daily_entry_service.process_submission(source_members, score_form(["Import_A", "Import_B"], 4), "2026-07-14")
    store.daily_entry_service.process_submission(source_members, score_form(["Import_A", "Import_B"], 6), "2026-07-15")
    store._context.export_service.export_to_excel()

    from bm2.services.store_application import StoreApplication

    target = StoreApplication(test_workspace / "import_target")
    shutil.copy2(store.workbook_path, target.workbook_path)

    target.refresh_local_database()
    first_import = _score_rows_for_names(target, {"Import_A", "Import_B"})
    target.refresh_local_database()
    second_import = _score_rows_for_names(target, {"Import_A", "Import_B"})

    assert len(first_import) == 4
    assert len(second_import) == 4
    assert {
        (row["member_name"], row["score_date"], row["score"])
        for row in second_import
    } == {
        (row["member_name"], row["score_date"], row["score"])
        for row in first_import
    }


def test_empty_or_missing_column_excel_import_does_not_pollute_sqlite(store):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SCORE_SHEET
    sheet.cell(1, 1, "not a date")
    sheet.cell(1, 2, TOTAL_HEADER)
    # NAME_HEADER intentionally omitted.
    workbook.save(store.workbook_path)
    workbook.close()

    before_scores = store._context.local_db.get_score_rows(include_deleted=True)
    store.refresh_local_database()

    assert store._context.local_db.get_score_rows(include_deleted=True) == before_scores


def test_invalid_date_excel_import_does_not_create_scores(store):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SCORE_SHEET
    sheet.cell(1, 1, "invalid-date")
    sheet.cell(1, 2, TOTAL_HEADER)
    sheet.cell(1, 3, NAME_HEADER)
    sheet.cell(2, 1, 5)
    sheet.cell(2, 3, "Bad_Date_Member")
    workbook.save(store.workbook_path)
    workbook.close()

    store.refresh_local_database()

    assert _score_rows_for_names(store, {"Bad_Date_Member"}) == []
