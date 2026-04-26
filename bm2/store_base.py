from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .constants import (
    DATA_FILE_PATTERNS,
    DEFAULT_MEMBERS,
    DEFAULT_QUICK_SCORES,
    DISABLED,
    ENABLED,
    EXPENSE_NAME_HEADER,
    INCOME_NAME_HEADER,
    NAME_HEADER,
    SCORE_SHEET,
    TOTAL_HEADER,
    WEAR_ABNORMAL_THRESHOLD,
    WEAR_NAME_HEADER,
    WINDOW_SIZE,
    WORKBOOK_FILENAME_PREFIX,
    normalize_status,
)
from .repositories import ConfigRepository
from .ui_text import MESSAGES


class StoreBaseMixin:
    def __init__(self, base_dir: Path) -> None:
        self.base_dir = base_dir
        self.config_repository = ConfigRepository(base_dir / 'system_config.json')
        self.config_repository.ensure_member_config(self._timestamp)
        self.workbook_path = self._resolve_workbook_path()
        self._ensure_workbook()

    def _timestamp(self) -> str:
        return datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    def _load_config(self) -> dict[str, Any]:
        return self.config_repository.load()

    def _save_config(self, config: dict[str, Any]) -> None:
        self.config_repository.save(config)

    def get_members(self) -> list[dict[str, str]]:
        self.config_repository.ensure_member_config(self._timestamp)
        return self.config_repository.load_members(normalize_status)

    def get_member(self, name: str) -> dict[str, str] | None:
        for member in self.get_members():
            if member['name'] == name:
                return member
        return None

    def _sync_member_visibility_in_workbook(self, workbook) -> None:
        self._ensure_score_sheet_structure(workbook)
        self._ensure_wear_sheet_structure(workbook)
        self._ensure_income_sheet_structure(workbook)
        self._ensure_expense_sheet_structure(workbook)

        score_sheet = self._score_sheet(workbook)
        wear_sheet = self._wear_sheet(workbook)
        income_sheet = self._income_sheet(workbook)
        expense_sheet = self._expense_sheet(workbook)
        score_name_col = self._find_column(score_sheet, NAME_HEADER)
        wear_name_col = self._find_column(wear_sheet, WEAR_NAME_HEADER)
        income_name_col = self._find_column(income_sheet, INCOME_NAME_HEADER)
        expense_name_col = self._find_column(expense_sheet, EXPENSE_NAME_HEADER)

        for member in self.get_members():
            should_hide = member['status'] == DISABLED
            if score_name_col is not None:
                self._set_member_row_hidden(score_sheet, score_name_col, member['name'], should_hide)
            if wear_name_col is not None:
                self._set_member_row_hidden(wear_sheet, wear_name_col, member['name'], should_hide)
            if income_name_col is not None:
                self._set_member_row_hidden(income_sheet, income_name_col, member['name'], should_hide)
            if expense_name_col is not None:
                self._set_member_row_hidden(expense_sheet, expense_name_col, member['name'], should_hide)


    def _delete_member_from_workbook(self, workbook, member_name: str) -> None:
        score_sheet = self._score_sheet(workbook)
        wear_sheet = self._wear_sheet(workbook)
        income_sheet = self._income_sheet(workbook)
        expense_sheet = self._expense_sheet(workbook)
        score_name_col = self._find_column(score_sheet, NAME_HEADER)
        wear_name_col = self._find_column(wear_sheet, WEAR_NAME_HEADER)
        income_name_col = self._find_column(income_sheet, INCOME_NAME_HEADER)
        expense_name_col = self._find_column(expense_sheet, EXPENSE_NAME_HEADER)

        if expense_name_col is not None:
            self._delete_member_row(expense_sheet, expense_name_col, member_name)
        if income_name_col is not None:
            self._delete_member_row(income_sheet, income_name_col, member_name)
        if wear_name_col is not None:
            self._delete_member_row(wear_sheet, wear_name_col, member_name)

        if score_name_col is not None:
            self._delete_member_row(score_sheet, score_name_col, member_name)

    def _save_member_workbook_change(self, change_func) -> None:
        workbook = self._open_workbook()
        try:
            change_func(workbook)
            self._save_workbook(workbook)
        finally:
            workbook.close()

    def sync_members_to_workbook(self) -> None:
        self._save_member_workbook_change(self._sync_member_visibility_in_workbook)

    def delete_member_from_workbook(self, member_name: str) -> None:
        self._save_member_workbook_change(lambda workbook: self._delete_member_from_workbook(workbook, member_name))

    def get_quick_scores(self) -> list[int]:
        config = self._load_config()
        return [int(value) for value in config.get('quick_scores', DEFAULT_QUICK_SCORES)]

    def get_default_score_date(self) -> str:
        config = self._load_config()
        default_date = str(config.get('default_score_date') or '').strip()
        if default_date:
            try:
                return datetime.strptime(default_date, '%Y-%m-%d').strftime('%Y-%m-%d')
            except ValueError:
                pass
        return datetime.now().strftime('%Y-%m-%d')

    def get_wear_abnormal_threshold(self) -> float:
        override_value = getattr(self, '_wear_abnormal_threshold_override', None)
        if override_value is not None:
            return float(override_value)
        config = self._load_config()
        raw_value = config.get('wear_abnormal_threshold', WEAR_ABNORMAL_THRESHOLD)
        try:
            return float(raw_value)
        except (TypeError, ValueError):
            return float(WEAR_ABNORMAL_THRESHOLD)

    def set_wear_abnormal_threshold(self, value_text: str) -> float:
        threshold = float(value_text.strip())
        self._wear_abnormal_threshold_override = threshold
        workbook = self._open_workbook()
        try:
            self._ensure_wear_sheet_structure(workbook)
            self._save_workbook(workbook)
        finally:
            workbook.close()
            if hasattr(self, '_wear_abnormal_threshold_override'):
                delattr(self, '_wear_abnormal_threshold_override')

        config = self._load_config()
        config['wear_abnormal_threshold'] = threshold
        self._save_config(config)
        return threshold

    def set_default_score_date(self, date_text: str) -> str:
        saved_date = datetime.strptime(date_text.strip(), '%Y-%m-%d').strftime('%Y-%m-%d')
        config = self._load_config()
        config['default_score_date'] = saved_date
        self._save_config(config)
        return saved_date

    def create_new_cycle(self, start_date_text: str) -> str:
        raw_text = (start_date_text or '').strip()
        if not raw_text:
            raise ValueError(MESSAGES['new_cycle_date_required'])
        try:
            start_date = datetime.strptime(raw_text, '%Y-%m-%d').date()
        except ValueError as exc:
            raise ValueError(MESSAGES['new_cycle_date_invalid']) from exc

        date_text = start_date.strftime('%Y-%m-%d')
        new_path = self.base_dir / f"{WORKBOOK_FILENAME_PREFIX}{date_text}.xlsx"
        counter = 2
        while new_path.exists():
            new_path = self.base_dir / f"{WORKBOOK_FILENAME_PREFIX}{date_text}_{counter}.xlsx"
            counter += 1

        workbook = Workbook()
        score_sheet = workbook.active
        score_sheet.title = SCORE_SHEET
        for column_index in range(1, WINDOW_SIZE + 1):
            header_date = start_date - timedelta(days=(WINDOW_SIZE - column_index))
            score_sheet.cell(1, column_index, header_date.strftime('%m-%d'))
        score_sheet.cell(1, WINDOW_SIZE + 1, TOTAL_HEADER)
        score_sheet.cell(1, WINDOW_SIZE + 2, NAME_HEADER)

        try:
            workbook.save(new_path)
        except PermissionError as exc:
            raise ValueError(MESSAGES['excel_busy'].format(filename=new_path.name)) from exc
        except OSError as exc:
            raise ValueError(MESSAGES['excel_save_failed'].format(filename=new_path.name)) from exc
        finally:
            workbook.close()

        previous_path = self.workbook_path
        self.workbook_path = new_path
        try:
            self._ensure_workbook()
        except Exception:
            self.workbook_path = previous_path
            try:
                new_path.unlink()
            except OSError:
                pass
            raise

        config = self._load_config()
        config['excel_filename'] = new_path.name
        config['default_score_date'] = date_text
        self._save_config(config)
        return new_path.name

    def _resolve_workbook_path(self) -> Path:
        config = self._load_config()
        configured = config.get('excel_filename')
        if configured:
            candidate = self.base_dir / configured
            if candidate.exists():
                return candidate

        existing_files: list[Path] = []
        for pattern in DATA_FILE_PATTERNS:
            existing_files.extend(self.base_dir.glob(pattern))
        existing_files = sorted(set(existing_files))
        target = existing_files[0] if existing_files else self.base_dir / f"{WORKBOOK_FILENAME_PREFIX}{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        config['excel_filename'] = target.name
        self._save_config(config)
        return target

    def _open_workbook(self):
        return load_workbook(self.workbook_path)

    def _save_workbook(self, workbook) -> None:
        try:
            workbook.save(self.workbook_path)
        except PermissionError as exc:
            raise ValueError(MESSAGES['excel_busy'].format(filename=self.workbook_path.name)) from exc
        except OSError as exc:
            raise ValueError(MESSAGES['excel_save_failed'].format(filename=self.workbook_path.name)) from exc
