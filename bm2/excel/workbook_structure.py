from __future__ import annotations

from openpyxl import Workbook, load_workbook

from ..constants import (
    EXPENSE_META_SHEET_ALIASES,
    EXPENSE_NAME_HEADER,
    DISABLED,
    INCOME_META_SHEET_ALIASES,
    INCOME_NAME_HEADER,
    META_HEADERS,
    META_SHEET,
    META_SHEET_ALIASES,
    NAME_HEADER,
    PROFIT_SHEET,
    WEAR_META_SHEET_ALIASES,
    WEAR_NAME_HEADER,
    WINDOW_SIZE,
)
from .profit_recalculator import recalculate_score_profits


class StoreStructureMixin:
    def _create_initial_score_sheet(self, workbook) -> None:
        self.score_sheet.create_initial_sheet(workbook)

    def _ensure_meta_sheets(self, workbook) -> bool:
        changed = False
        _, meta_changed = self._ensure_named_sheet(workbook, META_SHEET, META_SHEET_ALIASES, META_HEADERS, hidden=True)
        changed = meta_changed or changed
        for sheet_type in ('wear', 'income', 'expense'):
            _, value_meta_changed = self._ensure_named_sheet(
                workbook,
                self._value_sheet_spec(sheet_type)['meta_sheet'],
                self._value_sheet_spec(sheet_type)['meta_aliases'],
                self._value_sheet_spec(sheet_type)['meta_headers'],
                hidden=True,
            )
            changed = value_meta_changed or changed
        return changed

    def _ensure_score_sheet_structure(self, workbook) -> bool:
        return self.score_sheet.ensure_structure(workbook)

    def _recalculate_score_profits(self, workbook, score_sheet, profit_col: int) -> None:
        recalculate_score_profits(workbook, score_sheet, profit_col, self.income_sheet, self.wear_sheet)

    def _recalculate_wear_totals(self, sheet, total_col: int) -> None:
        self.wear_sheet.recalculate_totals(sheet, total_col)

    def _ensure_wear_sheet_structure(self, workbook) -> bool:
        return self.wear_sheet.ensure_structure(workbook)

    def _ensure_income_sheet_structure(self, workbook) -> bool:
        return self.income_sheet.ensure_structure(workbook)

    def _ensure_expense_sheet_structure(self, workbook) -> bool:
        return self.expense_sheet.ensure_structure(workbook)

    def _ensure_workbook(self) -> None:
        changed = False
        if self.workbook_path.exists():
            workbook = load_workbook(self.workbook_path)
        else:
            if getattr(self, 'read_only', False):
                raise FileNotFoundError(f"Workbook not found: {self.workbook_path}")
            workbook = Workbook()
            self._create_initial_score_sheet(workbook)
            changed = True

        try:
            changed = self._ensure_meta_sheets(workbook) or changed
            changed = self._ensure_score_sheet_structure(workbook) or changed
            changed = self._ensure_wear_sheet_structure(workbook) or changed
            changed = self._ensure_income_sheet_structure(workbook) or changed
            changed = self._ensure_expense_sheet_structure(workbook) or changed
            if PROFIT_SHEET in workbook.sheetnames:
                workbook.remove(workbook[PROFIT_SHEET])
                changed = True
            if not getattr(self, 'read_only', False):
                changed = self._sync_member_visibility_in_workbook(workbook) or changed
            if changed and not getattr(self, 'read_only', False):
                self.workbook_repository.save(workbook)
        finally:
            workbook.close()

    def _sync_member_visibility_in_workbook(self, workbook) -> bool:
        changed = False
        self._ensure_score_sheet_structure(workbook)
        self._ensure_wear_sheet_structure(workbook)
        self._ensure_income_sheet_structure(workbook)
        self._ensure_expense_sheet_structure(workbook)

        score_sheet = self._score_sheet(workbook)
        wear_sheet = self._wear_sheet(workbook)
        income_sheet = self._income_sheet(workbook)
        expense_sheet = self._expense_sheet(workbook)
        score_name_col = self._find_column(score_sheet, NAME_HEADER)
        wear_name_col = self._find_column(wear_sheet, WEAR_NAME_HEADER)
        income_name_col = self._find_column(income_sheet, INCOME_NAME_HEADER)
        expense_name_col = self._find_column(expense_sheet, EXPENSE_NAME_HEADER)

        for member in self.get_members():
            should_hide = member['status'] == DISABLED
            if score_name_col is not None:
                changed = self._set_member_row_hidden(score_sheet, score_name_col, member['name'], should_hide) or changed
            if wear_name_col is not None:
                changed = self._set_member_row_hidden(wear_sheet, wear_name_col, member['name'], should_hide) or changed
            if income_name_col is not None:
                changed = self._set_member_row_hidden(income_sheet, income_name_col, member['name'], should_hide) or changed
            if expense_name_col is not None:
                changed = self._set_member_row_hidden(expense_sheet, expense_name_col, member['name'], should_hide) or changed
        return changed

    def _recalculate_totals(self, sheet, total_col: int) -> None:
        self.score_sheet.recalculate_totals(sheet, total_col)

    def _format_score_sheet(self, sheet, recent_numbers: list[int], total_col: int) -> None:
        self.score_sheet.format_sheet(sheet, recent_numbers, total_col)
