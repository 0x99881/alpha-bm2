from __future__ import annotations

from typing import Any

from openpyxl.styles import Font

from .header_locator import find_column, find_sheet_by_alias, numeric_header_columns
from .member_rows import ensure_member_rows, sort_named_rows
from .summary_columns import ensure_summary_columns
from .value_normalizer import normalize_wear
from ..constants import (
    DATA_START_ROW,
    WEAR_ABNORMAL_FONT_COLOR,
    WEAR_NAME_HEADER,
    WEAR_SHEET,
    WEAR_SHEET_ALIASES,
    WEAR_TOTAL_HEADER,
)
from ..ui_text import UI_TEXT
from ..value_utils import to_float_or_none


class WearSheet:
    def __init__(self, store) -> None:
        self.store = store

    @property
    def daily_average_label(self) -> str:
        return UI_TEXT['wear_daily_member_avg']

    def sheet(self, workbook):
        sheet = find_sheet_by_alias(workbook, WEAR_SHEET_ALIASES)
        if sheet is not None:
            if sheet.title != WEAR_SHEET:
                sheet.title = WEAR_SHEET
            return sheet
        return workbook.create_sheet(title=WEAR_SHEET)

    def columns(self, sheet) -> list[tuple[int, int]]:
        return numeric_header_columns(sheet)

    def recalculate_totals(self, sheet, total_col: int) -> None:
        name_col = find_column(sheet, WEAR_NAME_HEADER)
        if name_col is None:
            return
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            name_value = str(sheet.cell(row, name_col).value or '').strip()
            if not name_value or name_value == self.daily_average_label:
                continue
            total = 0.0
            for _, col in self.columns(sheet):
                value = sheet.cell(row, col).value
                numeric = normalize_wear(value or 0)
                sheet.cell(row, col, numeric)
                total += numeric
            sheet.cell(row, total_col, normalize_wear(total))

    def format_sheet(self, sheet, total_col: int) -> None:
        threshold = self.store.get_wear_abnormal_threshold()
        normal_font = Font(color='000000', bold=False)
        abnormal_font = Font(color=WEAR_ABNORMAL_FONT_COLOR, bold=True)
        for _, wear_col in self.columns(sheet):
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                cell = sheet.cell(row, wear_col)
                value = cell.value
                if value in (None, ''):
                    cell.font = normal_font
                    continue
                try:
                    numeric_value = float(value)
                except (TypeError, ValueError):
                    cell.font = normal_font
                    continue
                cell.font = abnormal_font if numeric_value > threshold else normal_font

        for row in range(DATA_START_ROW, sheet.max_row + 1):
            sheet.cell(row, total_col).font = normal_font

    def _daily_average_rows(self, sheet, name_col: int) -> list[int]:
        return [
            row
            for row in range(DATA_START_ROW, sheet.max_row + 1)
            if str(sheet.cell(row, name_col).value or '').strip() == self.daily_average_label
        ]

    def _remove_daily_average_rows(self, sheet, name_col: int) -> bool:
        changed = False
        for row in reversed(self._daily_average_rows(sheet, name_col)):
            sheet.delete_rows(row, 1)
            changed = True
        return changed

    def _data_rows(self, sheet, name_col: int) -> list[int]:
        rows = []
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            name_value = str(sheet.cell(row, name_col).value or '').strip()
            if not name_value or name_value == self.daily_average_label:
                continue
            if bool(sheet.row_dimensions[row].hidden):
                continue
            rows.append(row)
        return rows

    def _set_cell_value(self, sheet, row: int, col: int, value) -> bool:
        if sheet.cell(row, col).value == value:
            return False
        sheet.cell(row, col, value)
        return True

    def _write_daily_average_row(self, sheet, *, total_col: int, name_col: int) -> bool:
        changed = False
        summary_rows = self._daily_average_rows(sheet, name_col)
        if len(summary_rows) == 1 and summary_rows[0] == sheet.max_row:
            summary_row = summary_rows[0]
        else:
            changed = self._remove_daily_average_rows(sheet, name_col) or changed
            summary_row = sheet.max_row + 1
            changed = True

        data_rows = self._data_rows(sheet, name_col)
        averages: dict[int, float] = {}
        for _, col in self.columns(sheet):
            values = [
                numeric_value
                for row in data_rows
                if (numeric_value := to_float_or_none(sheet.cell(row, col).value)) is not None
            ]
            averages[col] = normalize_wear(sum(values) / len(values)) if values else 0.0

        target_values = {col: value for col, value in averages.items()}
        target_values[total_col] = normalize_wear(sum(averages.values()))
        target_values[name_col] = self.daily_average_label
        for col in range(1, sheet.max_column + 1):
            value = target_values.get(col, None)
            changed = self._set_cell_value(sheet, summary_row, col, value) or changed
            sheet.cell(summary_row, col).font = Font(color='000000', bold=True)
        return changed

    def ensure_structure(self, workbook) -> bool:
        sheet = self.sheet(workbook)
        changed = False

        total_col, name_col, tail_changed = ensure_summary_columns(sheet, WEAR_TOTAL_HEADER, WEAR_NAME_HEADER)
        changed = changed or tail_changed

        changed = ensure_member_rows(
            sheet,
            members=self.store.get_members(),
            name_col=name_col,
            total_col=total_col,
            value_columns=[col for _, col in self.columns(sheet)],
        ) or changed

        self.recalculate_totals(sheet, total_col)
        sort_named_rows(sheet, total_col, name_col)
        changed = self._write_daily_average_row(sheet, total_col=total_col, name_col=name_col) or changed
        self.format_sheet(sheet, total_col)
        return changed

    def snapshot(self, workbook) -> dict[str, Any]:
        self.ensure_structure(workbook)
        sheet = self.sheet(workbook)
        headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
        raw_rows = []
        name_col = find_column(sheet, WEAR_NAME_HEADER)
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            if bool(sheet.row_dimensions[row].hidden):
                continue
            values = [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)]
            if name_col is not None and not values[name_col - 1]:
                continue
            if name_col is not None and str(values[name_col - 1] or '').strip() == self.daily_average_label:
                continue
            raw_rows.append(values)
        wear_col_indices = [col - 1 for _, col in self.columns(sheet)]
        return {
            'headers': headers,
            'raw_rows': raw_rows,
            'wear_col_indices': wear_col_indices,
            'threshold': self.store.get_wear_abnormal_threshold(),
        }
