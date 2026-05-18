"""Cycle-block layout for wear / income / expense value sheets.

Each sheet is rendered as a stack of per-cycle blocks. The active (latest)
cycle sits at the top so existing row-1 readers — find_column, numeric_header
columns, etc. — keep working without modification. Historical (settled) cycles
appear below as self-contained sub-tables with their own banner, headers,
member rows, totals, and averages.

The layout is regenerated from SQLite end-to-end on every export, so the Excel
file is always a clean projection of the database. No incremental column
inserts. No member-row deduping. Just a rewrite.
"""
from __future__ import annotations

from datetime import date as _date_cls, datetime, timedelta
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill

from ..constants import (
    DATA_START_ROW,
    EXPENSE_NAME_HEADER,
    INCOME_NAME_HEADER,
    NAME_HEADER,
    RESERVED_MEMBER_NAMES,
    WEAR_NAME_HEADER,
    WEAR_TOTAL_HEADER,
)
from ..domain.rules.daily_entry import IncompleteBalanceInput, resolve_wear_value
from ..value_utils import to_float_or_none
from .sheet_metadata import replace_sheet_meta
from .value_normalizer import normalize_expense, normalize_income, normalize_wear, parse_decimal
from .value_sheet_spec import get_value_sheet_spec


CYCLE_BANNER_PREFIX = "═══ 周期"
CYCLE_TOTAL_LABEL = "周期合计"
CYCLE_AVG_LABEL = "周期平均"
WEAR_DAILY_AVG_LABEL = "单日平均单号磨损"

_BANNER_FILL = PatternFill(start_color="FFEEE8", end_color="FFEEE8", fill_type="solid")
_SUMMARY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")


_NORMALIZERS = {
    "wear": normalize_wear,
    "income": normalize_income,
    "expense": normalize_expense,
}


def _text(value: Any) -> str:
    return str(value or "").strip()


def _day_code(date_iso: str) -> str:
    """``2026-05-12`` -> ``0512`` (the existing 4-digit header format)."""
    return _text(date_iso)[5:].replace("-", "")


def _resolve_wear(row: dict[str, Any]) -> float | None:
    name = _text(row.get("member_name"))
    try:
        value = resolve_wear_value(
            before_text=_text(row.get("before_balance")),
            after_text=_text(row.get("after_balance")),
            manual_wear_text=_text(row.get("manual_wear")),
            parse_decimal=parse_decimal,
            manual_field=f"{name} 手动磨损",
            before_field=f"{name} 前余额",
            after_field=f"{name} 后余额",
        )
    except (IncompleteBalanceInput, ValueError):
        return None
    return None if value is None else normalize_wear(value)


def _resolve_value(row: dict[str, Any], sheet_type: str) -> float | None:
    if sheet_type == "wear":
        return _resolve_wear(row)
    key = "income" if sheet_type == "income" else "other_expense"
    raw = _text(row.get(key))
    if not raw:
        return None
    parsed = to_float_or_none(raw)
    if parsed is None:
        return None
    normalizer = _NORMALIZERS[sheet_type]
    return normalizer(parsed)


def build_cycle_blocks(
    *, local_db, sheet_type: str, member_order: list[str],
) -> list[dict[str, Any]]:
    """Return cycle blocks for ``sheet_type``, newest-first.

    Each block dict carries: ``cycle`` (raw db row), ``dates`` (ISO list),
    ``member_values`` (``{name: {date: value}}``), ``is_settled``,
    ``is_active`` (the top block).

    Fallback: when no settlement cycle exists yet but score entries do,
    return a single synthetic "current period" block so refresh-from-Excel
    and snapshot views still work on a fresh install or pre-cycle workbook.
    """
    cycles = local_db.get_settlement_cycles()
    score_rows = local_db.get_score_rows()

    by_date: dict[str, list[dict[str, Any]]] = {}
    for row in score_rows:
        date_text = _text(row.get("score_date"))
        if date_text:
            by_date.setdefault(date_text, []).append(row)

    if not cycles:
        if not by_date:
            return []
        # Synthetic block: every entry is active, no cycle metadata.
        return [_synthetic_block(by_date, sheet_type, member_order)]

    def _sort_key(cycle: dict[str, Any]) -> str:
        return _text(cycle.get("start_date")) or _text(cycle.get("created_at"))

    cycles_chronological = sorted(cycles, key=_sort_key)
    cycles_newest_first = list(reversed(cycles_chronological))

    today_iso = _date_cls.today().strftime("%Y-%m-%d")
    blocks: list[dict[str, Any]] = []
    for index, cycle in enumerate(cycles_newest_first):
        start = _text(cycle.get("start_date"))
        if not start:
            continue
        is_settled = bool(int(cycle.get("settled", 0) or 0))
        settle = _text(cycle.get("settle_date"))
        # Active window ends today; settled window ends at settle_date.
        end_inclusive = settle if (is_settled and settle) else today_iso

        member_values: dict[str, dict[str, float]] = {name: {} for name in member_order}
        # Every day in [start, end_inclusive] becomes a column, even if no
        # member had a value that day. Keeps the table rectangular so users
        # can see "this day existed but everyone was 0" at a glance instead
        # of silently collapsing days out of the layout.
        date_list = _enumerate_dates(start, end_inclusive)
        for date_text in date_list:
            rows = by_date.get(date_text, [])
            for row in rows:
                name = _text(row.get("member_name"))
                if not name:
                    continue
                value = _resolve_value(row, sheet_type)
                if value is None:
                    continue
                member_values.setdefault(name, {})[date_text] = value

        blocks.append({
            "cycle": cycle,
            "dates": date_list,
            "member_values": member_values,
            "is_settled": is_settled,
            "is_active": index == 0,  # newest is active block
        })
    return blocks


def _enumerate_dates(start_iso: str, end_iso: str) -> list[str]:
    """Inclusive list of ISO dates from ``start_iso`` to ``end_iso``."""
    try:
        start = datetime.strptime(start_iso, "%Y-%m-%d").date()
        end = datetime.strptime(end_iso, "%Y-%m-%d").date()
    except ValueError:
        return []
    if end < start:
        return []
    out: list[str] = []
    cursor = start
    while cursor <= end:
        out.append(cursor.strftime("%Y-%m-%d"))
        cursor += timedelta(days=1)
    return out


_BLOCK_TERMINATORS = set(RESERVED_MEMBER_NAMES)


def _parse_banner_dates(banner_text: str) -> tuple[str, str]:
    """Extract (start_iso, end_iso) from a banner like
    ``═══ 周期 名字  2026-05-10 ~ 2026-05-15  ═══`` (Chinese full-width
    spaces / regular spaces both accepted). End may be ``进行中`` for
    open cycles, in which case ``end_iso`` is returned as empty.
    """
    import re

    matches = re.findall(r"(\d{4}-\d{2}-\d{2})", banner_text or "")
    if not matches:
        return "", ""
    start_iso = matches[0]
    end_iso = matches[1] if len(matches) >= 2 else ""
    return start_iso, end_iso


def _mmdd_to_iso(mmdd: str, start_iso: str, end_iso: str, today_iso: str) -> str:
    """Resolve a 4-digit MMDD column header to an ISO date.

    Within a cycle window: use the start_iso's year; bump to start_year+1
    if the resulting date falls outside the window (handles year wrap-
    over). Without a window (active block, no banner): use the current
    year, falling back to previous year if the date would be in the
    future.
    """
    text = str(mmdd or "").strip()
    if len(text) != 4 or not text.isdigit():
        return ""
    month = int(text[:2])
    day = int(text[2:])
    if start_iso:
        try:
            start_year = int(start_iso[:4])
        except ValueError:
            start_year = datetime.now().year
        for year_offset in (0, 1):
            try:
                cand = _date_cls(start_year + year_offset, month, day).strftime("%Y-%m-%d")
            except ValueError:
                continue
            if start_iso <= cand and (not end_iso or cand <= end_iso):
                return cand
    # Active-block fallback when no banner is available.
    year = int(today_iso[:4]) if today_iso else datetime.now().year
    for year_offset in (0, -1):
        try:
            cand = _date_cls(year + year_offset, month, day).strftime("%Y-%m-%d")
        except ValueError:
            continue
        if cand <= (today_iso or cand):
            return cand
    return ""


def _looks_like_summary_row(sheet, row: int, name_col: int) -> bool:
    name_cell = sheet.cell(row, name_col)
    if not bool(getattr(name_cell.font, "bold", False)):
        return False
    for col in range(1, name_col):
        if bool(getattr(sheet.cell(row, col).font, "bold", False)):
            return True
    return False


def iter_blocks(sheet, name_header: str, total_header: str | None) -> list[dict[str, Any]]:
    """Walk a value sheet rendered by ``rewrite_value_sheet_with_blocks``
    and yield each block's structure.

    Each entry has:
      - ``banner_row``: int or None (None for the active block on row 1)
      - ``start_iso`` / ``end_iso``: from the banner (both '' for active)
      - ``header_row``: int (the row with date / total / name headers)
      - ``name_col``: int — name column within this block
      - ``date_by_col``: {col -> ISO date}
      - ``data_rows``: list of int (member rows)
    Returns [] if the sheet doesn't appear to use the block layout.
    """
    max_row = sheet.max_row
    max_col = sheet.max_column
    today_iso = _date_cls.today().strftime("%Y-%m-%d")

    # Locate all banner rows (historical blocks).
    banner_rows: list[int] = []
    for row in range(1, max_row + 1):
        first = str(sheet.cell(row, 1).value or "")
        if first.startswith(CYCLE_BANNER_PREFIX):
            banner_rows.append(row)

    # Active block sits on row 1 (header). Historical blocks start at
    # banner_row + 1. block_headers is a list of (banner_row|None,
    # header_row, block_end_inclusive).
    block_starts: list[tuple[int | None, int]] = [(None, 1)] + [(br, br + 1) for br in banner_rows]

    blocks: list[dict[str, Any]] = []
    for index, (banner_row, header_row) in enumerate(block_starts):
        if index + 1 < len(block_starts):
            block_end = block_starts[index + 1][0] - 1  # one before next banner
        else:
            block_end = max_row

        # Resolve the cycle window from the banner.
        if banner_row is not None:
            banner_text = str(sheet.cell(banner_row, 1).value or "")
            start_iso, end_iso = _parse_banner_dates(banner_text)
        else:
            start_iso, end_iso = "", ""

        # Find the name column within THIS block by scanning the header row.
        name_col = None
        for col in range(1, max_col + 1):
            if str(sheet.cell(header_row, col).value or "") == name_header:
                name_col = col
                break
        if name_col is None:
            continue

        # Date columns: numeric MMDD headers in the header row, before name_col,
        # excluding the total_header column.
        date_by_col: dict[int, str] = {}
        unresolved_headers: list[str] = []
        for col in range(1, name_col):
            raw = sheet.cell(header_row, col).value
            if total_header is not None and raw == total_header:
                continue
            mmdd = str(raw or "").strip()
            iso = _mmdd_to_iso(mmdd, start_iso, end_iso, today_iso)
            if iso:
                date_by_col[col] = iso
            elif len(mmdd) == 4 and mmdd.isdigit():
                unresolved_headers.append(mmdd)
        if unresolved_headers:
            from ..ui_text import MESSAGES

            raise ValueError(
                MESSAGES["excel_cycle_block_date_unresolved"].format(
                    columns=", ".join(unresolved_headers)
                )
            )

        # Data rows: rows past the header, until the block end, skipping
        # blank-name rows and summary terminator rows.
        data_rows: list[int] = []
        for row in range(header_row + 1, block_end + 1):
            if bool(sheet.row_dimensions[row].hidden):
                continue
            name_val = str(sheet.cell(row, name_col).value or "").strip()
            if not name_val:
                continue
            if name_val in _BLOCK_TERMINATORS and _looks_like_summary_row(sheet, row, name_col):
                continue
            data_rows.append(row)

        blocks.append({
            "banner_row": banner_row,
            "start_iso": start_iso,
            "end_iso": end_iso,
            "header_row": header_row,
            "name_col": name_col,
            "date_by_col": date_by_col,
            "data_rows": data_rows,
        })
    return blocks


def _synthetic_block(
    by_date: dict[str, list[dict[str, Any]]],
    sheet_type: str,
    member_order: list[str],
) -> dict[str, Any]:
    member_values: dict[str, dict[str, float]] = {name: {} for name in member_order}
    date_set: set[str] = set()
    for date_text, rows in by_date.items():
        for row in rows:
            name = _text(row.get("member_name"))
            if not name:
                continue
            value = _resolve_value(row, sheet_type)
            if value is None:
                continue
            member_values.setdefault(name, {})[date_text] = value
            date_set.add(date_text)
    return {
        "cycle": {"id": "", "name": "当前", "start_date": "", "settle_date": ""},
        "dates": sorted(date_set),
        "member_values": member_values,
        "is_settled": False,
        "is_active": True,
    }


def _banner_text(block: dict[str, Any]) -> str:
    cycle = block["cycle"]
    name = _text(cycle.get("name")) or "周期"
    start = _text(cycle.get("start_date"))
    end = _text(cycle.get("settle_date")) or "进行中"
    return f"{CYCLE_BANNER_PREFIX} {name}　{start} ~ {end}　═══"


def _write_block(
    sheet,
    *,
    sheet_type: str,
    block: dict[str, Any],
    member_order: list[str],
    start_row: int,
    name_header: str,
    total_header: str | None,
) -> int:
    """Write one block starting at ``start_row``. Returns next free row.

    Column layout matches the legacy 2D table exactly so the active block
    drops in on row 1 without breaking any existing reader:
        ``[date_1] [date_2] ... [date_N] [total_header? (累计磨损)] [name_header]``

    Cycle-level summaries (合计 / 平均) appear as **rows** below the member
    rows, not columns.
    """
    dates = list(block["dates"])
    is_wear = sheet_type == "wear"

    # ---- header row -----
    header_row = start_row
    for col_index, date_iso in enumerate(dates, start=1):
        sheet.cell(header_row, col_index, _day_code(date_iso))
    next_col = len(dates) + 1
    total_col = None
    if total_header:
        total_col = next_col
        sheet.cell(header_row, total_col, total_header)
        next_col += 1
    name_col = next_col
    sheet.cell(header_row, name_col, name_header)
    for col in range(1, name_col + 1):
        sheet.cell(header_row, col).font = Font(bold=True)

    # ---- member rows -----
    if block["is_active"]:
        members_to_render = list(member_order)
    else:
        members_to_render = [name for name in member_order if block["member_values"].get(name)]

    member_rows: list[int] = []
    for offset, name in enumerate(members_to_render, start=1):
        row = header_row + offset
        member_rows.append(row)
        member_values = block["member_values"].get(name, {})
        per_date_sum = 0.0
        for col_index, date_iso in enumerate(dates, start=1):
            value = member_values.get(date_iso)
            if value is not None:
                sheet.cell(row, col_index, value)
                per_date_sum += float(value)
        if total_col is not None:
            sheet.cell(row, total_col, _NORMALIZERS[sheet_type](per_date_sum))
        sheet.cell(row, name_col, name)

    if not member_rows:
        return header_row + 1

    # ---- summary rows -----
    next_row = member_rows[-1] + 1
    daily_totals: dict[int, float] = {}
    for col_index, _ in enumerate(dates, start=1):
        s = 0.0
        for row in member_rows:
            cell_value = sheet.cell(row, col_index).value
            if isinstance(cell_value, (int, float)):
                s += float(cell_value)
        daily_totals[col_index] = s
    grand_total = sum(daily_totals.values())

    if block["is_active"]:
        # Active block: match the legacy structure — wear sheet has a 单日
        # 平均单号磨损 row; income/expense have no summary row. Cycle totals
        # are intentionally omitted here because they're shown on the page
        # cards / cycle profit view, and an extra row would confuse the
        # readers that scan all rows.
        if is_wear:
            # Per-column divisor: count only members who actually have a
            # value in THIS column. Using a single divisor (members with any
            # data) under-counts the per-day average when participation is
            # uneven across days.
            per_col_counts: dict[int, int] = {}
            for col_index in daily_totals.keys():
                per_col_counts[col_index] = sum(
                    1 for row in member_rows
                    if isinstance(sheet.cell(row, col_index).value, (int, float))
                )
            avg_row = next_row
            for col_index, value in daily_totals.items():
                divisor = per_col_counts.get(col_index, 0)
                sheet.cell(
                    avg_row, col_index,
                    _NORMALIZERS[sheet_type](value / divisor) if divisor else 0.0,
                )
            if total_col is not None:
                # The 累计磨损 column average is members' per-row totals
                # averaged across members who have ANY data this block.
                members_with_any = sum(
                    1 for row in member_rows
                    if any(isinstance(sheet.cell(row, c).value, (int, float)) for c in range(1, len(dates) + 1))
                )
                sheet.cell(
                    avg_row, total_col,
                    _NORMALIZERS[sheet_type](grand_total / members_with_any) if members_with_any else 0.0,
                )
            sheet.cell(avg_row, name_col, WEAR_DAILY_AVG_LABEL)
            for col in range(1, name_col + 1):
                cell = sheet.cell(avg_row, col)
                cell.font = Font(bold=True)
            return avg_row + 1
        return next_row

    # Historical block: full cycle summary — 合计 + 平均 rows.
    total_row = next_row
    for col_index, value in daily_totals.items():
        sheet.cell(total_row, col_index, _NORMALIZERS[sheet_type](value))
    if total_col is not None:
        sheet.cell(total_row, total_col, _NORMALIZERS[sheet_type](grand_total))
    sheet.cell(total_row, name_col, CYCLE_TOTAL_LABEL)
    for col in range(1, name_col + 1):
        cell = sheet.cell(total_row, col)
        cell.font = Font(bold=True)
        cell.fill = _SUMMARY_FILL

    avg_row = total_row + 1
    # Per-column divisor so uneven participation per day doesn't dilute the
    # average. Same rule as the active block's 单日平均 row.
    for col_index, value in daily_totals.items():
        col_count = sum(
            1 for row in member_rows
            if isinstance(sheet.cell(row, col_index).value, (int, float))
        )
        sheet.cell(
            avg_row, col_index,
            _NORMALIZERS[sheet_type](value / col_count) if col_count else 0.0,
        )
    if total_col is not None:
        members_with_any = sum(
            1 for row in member_rows
            if any(isinstance(sheet.cell(row, c).value, (int, float)) for c in range(1, len(dates) + 1))
        ) or len(member_rows) or 1
        sheet.cell(
            avg_row, total_col,
            _NORMALIZERS[sheet_type](grand_total / members_with_any),
        )
    sheet.cell(avg_row, name_col, CYCLE_AVG_LABEL)
    for col in range(1, name_col + 1):
        cell = sheet.cell(avg_row, col)
        cell.font = Font(bold=True)
        cell.fill = _SUMMARY_FILL

    return avg_row + 1


def _wipe_sheet(sheet) -> None:
    """Truly clear a sheet — openpyxl's delete_rows/delete_cols can leave
    stale cells in some workbook states. Explicitly null every cell then
    delete the structural rows to keep the sheet small."""
    max_row = sheet.max_row
    max_col = sheet.max_column
    if max_row >= 1 and max_col >= 1:
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                sheet.cell(row, col).value = None
    if max_row >= 1:
        sheet.delete_rows(1, max_row)
    if max_col >= 1:
        sheet.delete_cols(1, max_col)
    for row_dim in list(sheet.row_dimensions.keys()):
        sheet.row_dimensions[row_dim].hidden = False
    for merged_range in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged_range))


def rewrite_value_sheet_with_blocks(
    workbook,
    *,
    sheet_type: str,
    sheet,
    blocks: list[dict[str, Any]],
    member_order: list[str],
) -> None:
    """Replace ``sheet``'s contents with the block layout for ``blocks``.

    ``blocks`` must be ordered newest-first. The first block becomes the top
    block whose header row sits on row 1, preserving compatibility with all
    readers that locate headers via ``find_column`` on row 1.

    Also rewrites the corresponding meta sheet so date headers can still be
    resolved to ISO dates by readers that consult the meta sheet (e.g. the
    calendar's per-month reader path, when it's not yet switched to DB).
    """
    spec = get_value_sheet_spec(sheet_type)
    name_header = spec["name_header"]
    total_header = spec["total_header"]

    _wipe_sheet(sheet)

    if not blocks:
        # Even with no cycle, leave a minimal usable header row so the page
        # doesn't crash on a fresh workbook.
        sheet.cell(1, 1, name_header)
        replace_sheet_meta(workbook, spec["meta_sheet"], [])
        return

    cur_row = 1
    # Active block (no banner above so its header lands on row 1).
    cur_row = _write_block(
        sheet,
        sheet_type=sheet_type,
        block=blocks[0],
        member_order=member_order,
        start_row=cur_row,
        name_header=name_header,
        total_header=total_header,
    )

    # Historical blocks below — each preceded by a banner.
    for block in blocks[1:]:
        cur_row += 1  # blank divider
        banner_row = cur_row
        sheet.cell(banner_row, 1, _banner_text(block))
        cell = sheet.cell(banner_row, 1)
        cell.font = Font(bold=True)
        cell.fill = _BANNER_FILL
        cell.alignment = Alignment(horizontal="left")
        cur_row = banner_row + 1
        cur_row = _write_block(
            sheet,
            sheet_type=sheet_type,
            block=block,
            member_order=member_order,
            start_row=cur_row,
            name_header=name_header,
            total_header=total_header,
        )

    # Meta sheet: only the active block's dates need entries (readers only
    # consult row 1, which is the active block).
    active_block = blocks[0]
    meta_rows = [(date_iso, _day_code(date_iso)) for date_iso in active_block["dates"]]
    replace_sheet_meta(workbook, spec["meta_sheet"], meta_rows)
