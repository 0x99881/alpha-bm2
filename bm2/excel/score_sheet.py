from __future__ import annotations

import re
from datetime import datetime, timedelta
from typing import Any

from openpyxl.styles import Font, PatternFill

from .daily_target import DailyTarget
from .header_locator import find_column, find_sheet_by_alias
from .member_rows import ensure_member_rows, sort_named_rows
from .profit_recalculator import recalculate_score_profits
from .sheet_metadata import parse_saved_date, read_sheet_meta
from .value_normalizer import normalize_income
from ..constants import (
    DATA_START_ROW,
    INITIAL_SCORE_HEADER_COUNT,
    NAME_HEADER,
    OLD_SCORE_COLUMN_FILL,
    PROFIT_HEADER,
    SCORE_SHEET,
    SCORE_SHEET_ALIASES,
    SCORE_LOW_DAILY_FONT_COLOR,
    SCORE_TOTAL_DARK_RGB,
    SCORE_TOTAL_FLAT_FILL,
    SCORE_TOTAL_LIGHT_RGB,
    TOTAL_HEADER,
    WINDOW_SIZE,
    META_SHEET,
)
from ..ui_text import MESSAGES


class ScoreSheet:
    def __init__(self, store) -> None:
        self.store = store

    def sheet(self, workbook):
        sheet = find_sheet_by_alias(workbook, SCORE_SHEET_ALIASES)
        if sheet is not None:
            if sheet.title != SCORE_SHEET:
                sheet.title = SCORE_SHEET
            return sheet
        if workbook.sheetnames:
            workbook[workbook.sheetnames[0]].title = SCORE_SHEET
            return workbook[SCORE_SHEET]
        return workbook.create_sheet(title=SCORE_SHEET)

    def create_initial_sheet(self, workbook) -> None:
        score_sheet = workbook.active
        score_sheet.title = SCORE_SHEET
        end_date = datetime.now().date()
        for column_index in range(1, INITIAL_SCORE_HEADER_COUNT + 1):
            header_date = end_date - timedelta(days=(INITIAL_SCORE_HEADER_COUNT - column_index))
            score_sheet.cell(1, column_index, header_date.strftime('%m-%d'))
        score_sheet.cell(1, INITIAL_SCORE_HEADER_COUNT + 1, TOTAL_HEADER)
        score_sheet.cell(1, INITIAL_SCORE_HEADER_COUNT + 2, NAME_HEADER)

    def is_date_header(self, value: Any) -> bool:
        return isinstance(value, str) and re.fullmatch(r"\d{2}-\d{2}", value) is not None

    def date_columns(self, sheet) -> list[tuple[int, int]]:
        total_col = find_column(sheet, TOTAL_HEADER)
        result = []
        if total_col is None:
            return result
        for number, col in enumerate(range(1, total_col), start=1):
            if sheet.cell(1, col).value is not None:
                result.append((number, col))
        return result

    def ensure_tail_columns(self, sheet) -> tuple[int, int, int, bool]:
        total_col = find_column(sheet, TOTAL_HEADER)
        profit_col = find_column(sheet, PROFIT_HEADER)
        name_col = find_column(sheet, NAME_HEADER)
        changed = False

        if total_col is not None and profit_col is not None and name_col is not None and profit_col == total_col + 1 and name_col == profit_col + 1:
            sheet.cell(1, total_col, TOTAL_HEADER)
            sheet.cell(1, profit_col, PROFIT_HEADER)
            sheet.cell(1, name_col, NAME_HEADER)
            return total_col, profit_col, name_col, changed

        total_values = [sheet.cell(row, total_col).value for row in range(1, sheet.max_row + 1)] if total_col is not None else [None] * sheet.max_row
        profit_values = [sheet.cell(row, profit_col).value for row in range(1, sheet.max_row + 1)] if profit_col is not None else [None] * sheet.max_row
        name_values = [sheet.cell(row, name_col).value for row in range(1, sheet.max_row + 1)] if name_col is not None else [None] * sheet.max_row

        existing_cols = [col for col in [total_col, profit_col, name_col] if col is not None]
        for col_index in sorted(existing_cols, reverse=True):
            sheet.delete_cols(col_index)
        insert_at = sheet.max_column + 1
        sheet.insert_cols(insert_at, 3)
        for row, value in enumerate(total_values, start=1):
            sheet.cell(row, insert_at, value)
        for row, value in enumerate(profit_values, start=1):
            sheet.cell(row, insert_at + 1, value)
        for row, value in enumerate(name_values, start=1):
            sheet.cell(row, insert_at + 2, value)
        sheet.cell(1, insert_at, TOTAL_HEADER)
        sheet.cell(1, insert_at + 1, PROFIT_HEADER)
        sheet.cell(1, insert_at + 2, NAME_HEADER)
        changed = True
        return insert_at, insert_at + 1, insert_at + 2, changed

    def ensure_structure(self, workbook) -> bool:
        sheet = self.sheet(workbook)
        changed = False

        col = 1
        while col <= sheet.max_column:
            if sheet.cell(1, col).value is not None:
                col += 1
                continue
            if any(sheet.cell(row, col).value is not None for row in range(DATA_START_ROW, sheet.max_row + 1)):
                col += 1
                continue
            sheet.delete_cols(col)
            changed = True

        total_col, profit_col, name_col, tail_changed = self.ensure_tail_columns(sheet)
        changed = changed or tail_changed

        changed = ensure_member_rows(
            sheet,
            members=self.store.get_members(),
            name_col=name_col,
            total_col=total_col,
            value_columns=[col for _, col in self.date_columns(sheet)],
            extra_columns=[profit_col],
        ) or changed

        self.recalculate_totals(sheet, total_col)
        recalculate_score_profits(workbook, sheet, profit_col, self.store.income_sheet, self.store.wear_sheet)
        sort_named_rows(sheet, total_col, name_col)
        recent_numbers = [number for number, _ in self.date_columns(sheet)][-WINDOW_SIZE:]
        self.format_sheet(sheet, recent_numbers, total_col)
        return changed

    def read_rankings(self, sheet) -> list[dict[str, int | str]]:
        total_col = find_column(sheet, TOTAL_HEADER)
        name_col = find_column(sheet, NAME_HEADER)
        rows = []
        if total_col and name_col:
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                if bool(sheet.row_dimensions[row].hidden):
                    continue
                name = sheet.cell(row, name_col).value
                total = sheet.cell(row, total_col).value
                if name:
                    rows.append({'name': str(name), 'total': int(total or 0)})
        rows.sort(key=lambda item: (-int(item['total']), str(item['name'])))
        return rows

    def latest_column(self, sheet) -> str:
        score_columns = self.date_columns(sheet)
        return str(sheet.cell(1, score_columns[-1][1]).value) if score_columns else '-'

    def column_count(self, sheet) -> int:
        return len(self.date_columns(sheet))

    def snapshot(self, workbook) -> dict[str, Any]:
        self.ensure_structure(workbook)
        sheet = self.sheet(workbook)
        date_columns = self.date_columns(sheet)
        recent_date_cols = [col for _, col in date_columns[-WINDOW_SIZE:]]
        total_col = find_column(sheet, TOTAL_HEADER)
        profit_col = find_column(sheet, PROFIT_HEADER)
        name_col = find_column(sheet, NAME_HEADER)
        visible_columns = list(recent_date_cols)
        for col in (total_col, profit_col, name_col):
            if col is not None:
                visible_columns.append(col)
        headers = [sheet.cell(1, col).value for col in visible_columns]
        raw_rows = []
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            if bool(sheet.row_dimensions[row].hidden):
                continue
            if name_col is not None and not sheet.cell(row, name_col).value:
                continue
            values = [sheet.cell(row, col).value for col in visible_columns]
            raw_rows.append(values)
        return {
            'headers': headers,
            'raw_rows': raw_rows,
        }

    def active_member_profit_map(self, sheet, active_members: list[dict[str, str]]) -> dict[str, float]:
        name_col = find_column(sheet, NAME_HEADER)
        profit_col = find_column(sheet, PROFIT_HEADER)
        profit_map = {member['name']: 0.0 for member in active_members}
        if name_col is not None and profit_col is not None:
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                member_name = str(sheet.cell(row, name_col).value or '').strip()
                if member_name not in profit_map:
                    continue
                profit_map[member_name] = normalize_income(sheet.cell(row, profit_col).value or 0)
        return profit_map

    def recalculate_totals(self, sheet, total_col: int) -> None:
        name_col = find_column(sheet, NAME_HEADER)
        if name_col is None:
            return
        date_columns = self.date_columns(sheet)
        recent_numbers = {number for number, _ in date_columns[-WINDOW_SIZE:]}
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            if not sheet.cell(row, name_col).value:
                continue
            total = 0
            for number, date_col in date_columns:
                value = sheet.cell(row, date_col).value
                numeric = int(value or 0)
                sheet.cell(row, date_col, numeric)
                if number in recent_numbers:
                    total += numeric
            sheet.cell(row, total_col, total)

    def format_sheet(self, sheet, recent_numbers: list[int], total_col: int) -> None:
        name_col = find_column(sheet, NAME_HEADER)
        old_fill = PatternFill(start_color=OLD_SCORE_COLUMN_FILL, end_color=OLD_SCORE_COLUMN_FILL, fill_type='solid')
        normal_font = Font(color='000000', bold=False)
        low_score_font = Font(color=SCORE_LOW_DAILY_FONT_COLOR, bold=True)
        recent_set = set(recent_numbers)
        for number, date_col in self.date_columns(sheet):
            is_recent = number in recent_set
            fill = PatternFill(fill_type=None) if is_recent else old_fill
            for row in range(DATA_START_ROW, sheet.max_row + 1):
                cell = sheet.cell(row, date_col)
                cell.fill = fill
                if not is_recent:
                    cell.font = normal_font
                    continue
                try:
                    numeric_value = float(cell.value or 0)
                except (TypeError, ValueError):
                    cell.font = normal_font
                    continue
                if numeric_value < 10:
                    cell.font = low_score_font
                else:
                    cell.font = normal_font

        totals = []
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            if name_col is not None and not sheet.cell(row, name_col).value:
                continue
            value = sheet.cell(row, total_col).value
            if isinstance(value, (int, float)):
                totals.append(float(value))
        max_score = max(totals) if totals else 0.0
        min_score = min(totals) if totals else 0.0
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            if name_col is not None and not sheet.cell(row, name_col).value:
                continue
            value = sheet.cell(row, total_col).value
            score = float(value) if isinstance(value, (int, float)) else 0.0
            if max_score == min_score:
                color = SCORE_TOTAL_FLAT_FILL
            else:
                ratio = (score - min_score) / (max_score - min_score)
                red = int(SCORE_TOTAL_LIGHT_RGB[0] + (SCORE_TOTAL_DARK_RGB[0] - SCORE_TOTAL_LIGHT_RGB[0]) * ratio)
                green = int(SCORE_TOTAL_LIGHT_RGB[1] + (SCORE_TOTAL_DARK_RGB[1] - SCORE_TOTAL_LIGHT_RGB[1]) * ratio)
                blue = int(SCORE_TOTAL_LIGHT_RGB[2] + (SCORE_TOTAL_DARK_RGB[2] - SCORE_TOTAL_LIGHT_RGB[2]) * ratio)
                color = f'{red:02X}{green:02X}{blue:02X}'
            sheet.cell(row, total_col).fill = PatternFill(start_color=color, end_color=color, fill_type='solid')

    def find_empty_column(self, sheet, mmdd_text: str, total_col: int) -> int | None:
        for col in range(1, total_col):
            if str(sheet.cell(1, col).value or '').strip() == mmdd_text and not self.column_has_data(sheet, col):
                return col
        return None

    def prepare_target(self, score_sheet, effective_date_text: str):
        total_col = find_column(score_sheet, TOTAL_HEADER)
        name_col = find_column(score_sheet, NAME_HEADER)
        if total_col is None or name_col is None:
            raise ValueError(MESSAGES['score_sheet_invalid'])

        score_header = effective_date_text[5:]
        reuse_col = self.find_empty_column(score_sheet, score_header, total_col)
        date_columns = self.date_columns(score_sheet)
        if reuse_col is not None:
            next_number = next((number for number, col in date_columns if col == reuse_col), len(date_columns))
            return DailyTarget(
                sheet=score_sheet,
                target_col=reuse_col,
                total_col=total_col,
                profit_col=total_col + 1,
                name_col=total_col + 2,
                header=score_header,
                next_number=next_number,
            )

        next_number = date_columns[-1][0] + 1 if date_columns else 1
        score_sheet.insert_cols(total_col, 1)
        target_col = total_col
        score_sheet.cell(1, target_col, score_header)
        return DailyTarget(
            sheet=score_sheet,
            target_col=target_col,
            total_col=total_col + 1,
            profit_col=total_col + 2,
            name_col=total_col + 3,
            header=score_header,
            next_number=next_number,
        )

    def column_has_data(self, sheet, col: int) -> bool:
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            value = sheet.cell(row, col).value
            if value in (None, '', 0):
                continue
            try:
                if int(value) != 0:
                    return True
            except (TypeError, ValueError):
                return True
        return False

    def latest_used_date(self, workbook):
        latest_meta_dates = [
            parse_saved_date(saved_date)
            for saved_date, _ in read_sheet_meta(workbook, META_SHEET)
        ]
        year_hint = max(latest_meta_dates).year if latest_meta_dates else datetime.now().year
        latest_dates = list(latest_meta_dates)

        sheet = self.sheet(workbook)
        for _, col in reversed(self.date_columns(sheet)):
            if not self.column_has_data(sheet, col):
                continue
            latest_score_header = str(sheet.cell(1, col).value or '').strip()
            if self.is_date_header(latest_score_header):
                latest_dates.append(datetime.strptime(f'{year_hint}-{latest_score_header}', '%Y-%m-%d').date())
                break
        return max(latest_dates) if latest_dates else None
