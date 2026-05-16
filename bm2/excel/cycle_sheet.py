"""Stacks every settlement cycle as one block inside a single Excel sheet.

Layout (one sheet, name = `周期盈亏记录`) — flat single-table per block:

  [Block — 周期 yyyy-mm-dd ~ yyyy-mm-dd  (已结算 | 未结算)]
    姓名 | {start}\n期初余额 | {settle}\n期末余额 | {date1} | {date2} | …
        | 红包合计 | 磨损合计 | 收入合计 | 目前盈亏 | 利润 | 差异
    member rows
  [blank row]
  [next block…]

* 期初/期末余额 sit side-by-side at the front so the balance change is the
  first thing the eye lands on.
* 磨损合计 / 收入合计 / 目前盈亏 come straight from daily entries — visible
  pre-settle too (running estimate).
* 利润 = 期末 − 期初 − 红包合计 (balance method, after settle).
* 差异 = 利润 − 目前盈亏 (surfaces daily-entry vs. balance inaccuracies).
* Blocks chronological (oldest first; new cycles appended).
* On any change we rewrite the whole sheet and drop the legacy per-cycle
  tabs (周期MM-DD) from the previous one-tab-per-cycle implementation.
"""
from __future__ import annotations

from datetime import datetime
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..constants import NAME_HEADER


OVERVIEW_SHEET_NAME = "周期盈亏记录"

HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
PROFIT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
PROFIT_FLOW_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
DELTA_FILL = PatternFill(start_color="FFE5E5", end_color="FFE5E5", fill_type="solid")
WEAR_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
INCOME_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
TITLE_SETTLED_FILL = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
TITLE_UNSETTLED_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
UNSETTLED_FONT = Font(color="888888", italic=True)


def _zh_date(date_text: str) -> str:
    try:
        d = datetime.strptime(date_text, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return date_text
    return f"{d.month}月{d.day}日"


def _balance_header(date_text: str, fallback: str) -> str:
    zh = _zh_date(date_text)
    return f"{zh}余额" if zh and date_text else fallback


def _drop_legacy_per_cycle_sheets(workbook) -> None:
    """Remove old `周期MM-DD` tabs from the previous one-sheet-per-cycle design."""
    legacy = [
        name for name in workbook.sheetnames
        if name.startswith("周期") and name != OVERVIEW_SHEET_NAME
    ]
    for name in legacy:
        del workbook[name]


def regenerate_overview_sheet(workbook_repository, cycles_data: list[dict[str, Any]]) -> str:
    """Rebuild the entire 周期盈亏记录 sheet from the given cycles."""
    book = workbook_repository.open()
    try:
        write_overview_sheet(book, cycles_data)
        workbook_repository.save(book)
    finally:
        book.close()
    return OVERVIEW_SHEET_NAME


# Backward-compatible aliases (callers still using the old names).
def regenerate_cycle_sheet(workbook_repository, cycles_data: list[dict[str, Any]]) -> str:
    return regenerate_overview_sheet(workbook_repository, cycles_data)


def delete_cycle_sheet(workbook_repository, _start_date: str) -> str:
    """Compatibility shim — actual cleanup happens via the regen path now."""
    return OVERVIEW_SHEET_NAME


def write_overview_sheet(workbook, cycles_data: list[dict[str, Any]]) -> str:
    _drop_legacy_per_cycle_sheets(workbook)
    if OVERVIEW_SHEET_NAME in workbook.sheetnames:
        del workbook[OVERVIEW_SHEET_NAME]
    sheet = workbook.create_sheet(title=OVERVIEW_SHEET_NAME)

    if not cycles_data:
        sheet.cell(1, 1, "暂无周期").font = UNSETTLED_FONT
        sheet.freeze_panes = "A2"
        return OVERVIEW_SHEET_NAME

    current_row = 1
    max_width_map: dict[int, int] = {}
    for cycle_data in cycles_data:
        rows_used, col_widths = _write_one_block(sheet, current_row, cycle_data)
        for col_index, width in col_widths.items():
            max_width_map[col_index] = max(max_width_map.get(col_index, 0), width)
        current_row += rows_used + 1  # 1 blank separator between blocks

    for col_index, width in max_width_map.items():
        sheet.column_dimensions[get_column_letter(col_index)].width = max(width, 8)

    sheet.freeze_panes = "A2"
    return OVERVIEW_SHEET_NAME


def _write_one_block(sheet, start_row: int, cycle_data: dict[str, Any]) -> tuple[int, dict[int, int]]:
    cycle = cycle_data["selected_cycle"]
    start_date = str(cycle.get("start_date") or "").strip()
    settle_date = str(cycle.get("settle_date") or "").strip()
    is_settled = bool(int(cycle.get("settled", 0) or 0))
    redpacket_dates: list[str] = list(cycle_data.get("redpacket_dates") or [])
    rows: list[dict[str, Any]] = list(cycle_data.get("rows") or [])

    # Flat single-table layout. Multi-line balance headers carry the date.
    start_header = (f"{start_date}\n期初余额") if start_date else "期初余额"
    settle_header = (f"{settle_date}\n期末余额") if settle_date else "期末余额"

    headers: list[str] = [NAME_HEADER, start_header, settle_header]
    headers.extend(redpacket_dates)
    headers.extend(["红包合计", "磨损合计", "收入合计", "目前盈亏", "利润", "差异"])

    # column indices
    name_col = 1
    start_balance_col = 2
    end_balance_col = 3
    first_date_col = 4
    date_columns = {d: first_date_col + idx for idx, d in enumerate(redpacket_dates)}
    after_dates_col = first_date_col + len(redpacket_dates)
    redpacket_total_col = after_dates_col
    wear_total_col = after_dates_col + 1
    income_total_col = after_dates_col + 2
    flow_col = after_dates_col + 3
    profit_col = after_dates_col + 4
    delta_col = after_dates_col + 5
    total_cols = delta_col

    # ---- title row -------------------------------------------------------
    title_text = _title_for(start_date, settle_date, is_settled)
    title_cell = sheet.cell(start_row, 1, title_text)
    title_cell.font = Font(bold=True, size=12, color="1F2937")
    title_cell.fill = TITLE_SETTLED_FILL if is_settled else TITLE_UNSETTLED_FILL
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=total_cols)

    # ---- header row ------------------------------------------------------
    header_row = start_row + 1
    for col_index, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col_index, value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Slightly taller header row so 2-line balance labels fit cleanly.
    sheet.row_dimensions[header_row].height = 32

    # ---- member rows -----------------------------------------------------
    body_start = header_row + 1
    for row_index, row in enumerate(rows, start=body_start):
        name = row["member_name"]
        sheet.cell(row_index, name_col, name).alignment = Alignment(horizontal="left")

        start_balance = _coerce_number(row.get("start_balance"))
        if start_balance is not None:
            sheet.cell(row_index, start_balance_col, start_balance)

        end_balance = _coerce_number(row.get("end_balance"))
        if is_settled and end_balance is not None:
            sheet.cell(row_index, end_balance_col, end_balance)
        elif not is_settled:
            cell = sheet.cell(row_index, end_balance_col, "未结算")
            cell.font = UNSETTLED_FONT
            cell.alignment = Alignment(horizontal="center")

        for date_text, amount in (row.get("redpacket_by_date") or {}).items():
            col = date_columns.get(date_text)
            if col is None or not amount:
                continue
            sheet.cell(row_index, col, _coerce_number(amount))

        rp_total = row.get("redpacket_total") or 0
        rp_cell = sheet.cell(row_index, redpacket_total_col, rp_total if rp_total else None)
        if rp_total:
            rp_cell.fill = TOTAL_FILL

        # 磨损/收入/目前盈亏 are live from daily entries — show always.
        wear_total = row.get("wear_total") or 0
        wear_cell = sheet.cell(row_index, wear_total_col, wear_total if wear_total else None)
        if wear_total:
            wear_cell.fill = WEAR_FILL
        income_total = row.get("income_total") or 0
        income_cell = sheet.cell(row_index, income_total_col, income_total if income_total else None)
        if income_total:
            income_cell.fill = INCOME_FILL
        flow_value = row.get("profit_flow")
        if flow_value is not None:
            flow_cell = sheet.cell(row_index, flow_col, flow_value)
            flow_cell.fill = PROFIT_FLOW_FILL
            flow_cell.font = Font(bold=True)

        # 利润 and 差异 depend on settle balance.
        _write_profit_cell(sheet, row_index, profit_col, is_settled, row.get("profit"), PROFIT_FILL)
        _write_profit_cell(sheet, row_index, delta_col, is_settled, row.get("profit_delta"), DELTA_FILL, bold=False)

    rows_used = 2 + len(rows)  # title + header + members

    col_widths: dict[int, int] = {}
    col_widths[name_col] = 10
    col_widths[start_balance_col] = 13
    col_widths[end_balance_col] = 13
    for date in redpacket_dates:
        col_widths[date_columns[date]] = 12  # "2026-05-04"
    col_widths[redpacket_total_col] = 10
    col_widths[wear_total_col] = 10
    col_widths[income_total_col] = 10
    col_widths[flow_col] = 10
    col_widths[profit_col] = 10
    col_widths[delta_col] = 10

    return rows_used, col_widths


def _title_for(start_date: str, settle_date: str, is_settled: bool) -> str:
    start_label = start_date or "起始日期未填"
    if is_settled and settle_date:
        return f"周期 {start_label} ~ {settle_date}    已结算"
    return f"周期 {start_label} ~ 未结算"


def _write_profit_cell(sheet, row_index: int, col: int, is_settled: bool, value, fill, bold: bool = True) -> None:
    if is_settled and value is not None:
        cell = sheet.cell(row_index, col, value)
        cell.fill = fill
        if bold:
            cell.font = Font(bold=True)
    else:
        cell = sheet.cell(row_index, col, "未结算")
        cell.font = UNSETTLED_FONT
        cell.alignment = Alignment(horizontal="center")


def _coerce_number(value: Any):
    if value in (None, ""):
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        return value
    return int(f) if f.is_integer() else f
